"""
Prospect file parser — CSV and XLSX support with dynamic column mapping.

Two-phase upload flow:
  1. preview()  — reads headers + sample rows, returns them for column mapping UI
  2. parse()    — applies a column mapping and validates all rows into prospect dicts

Replaces the old csv_parser.py (kept for backwards compat via re-export).
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
from typing import Any, Dict, List, Optional, Tuple

from fastapi import UploadFile

# ---------------------------------------------------------------------------
# XLSX support — optional import so the module still loads if openpyxl
# isn't installed (e.g. during unit tests with minimal deps).
# ---------------------------------------------------------------------------
try:
    from openpyxl import load_workbook

    XLSX_AVAILABLE = True
except ImportError:
    XLSX_AVAILABLE = False
    load_workbook = None  # type: ignore[assignment,misc]


# ---------------------------------------------------------------------------
# Exceptions
# ---------------------------------------------------------------------------

class FileValidationError(Exception):
    """Raised when the uploaded file fails structural validation."""


# Keep the old name so existing import sites don't break.
CSVValidationError = FileValidationError


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

# System fields that uploaded columns can be mapped to.
SYSTEM_FIELDS = [
    "email",
    "first_name",
    "last_name",
    "company_name",
    "company_domain",
    "title",
    "phone",
    "linkedin_url",
    "industry",
    "company_size",
]

REQUIRED_SYSTEM_FIELDS = {"email"}

MAX_FILE_SIZE = 50 * 1024 * 1024  # 50 MB
MAX_ROWS = 100_000
PREVIEW_ROWS = 5  # rows returned in preview

EMAIL_RE = re.compile(r"^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$")

# Common header synonyms → system field.  Keys are lowercase / stripped.
HEADER_ALIASES: Dict[str, str] = {
    # email
    "email": "email",
    "email_address": "email",
    "e-mail": "email",
    "email address": "email",
    "emailaddress": "email",
    "work email": "email",
    "work_email": "email",
    "business email": "email",
    # first name
    "first_name": "first_name",
    "first name": "first_name",
    "firstname": "first_name",
    "fname": "first_name",
    "given name": "first_name",
    # last name
    "last_name": "last_name",
    "last name": "last_name",
    "lastname": "last_name",
    "lname": "last_name",
    "surname": "last_name",
    "family name": "last_name",
    # company
    "company_name": "company_name",
    "company name": "company_name",
    "company": "company_name",
    "organization": "company_name",
    "organisation": "company_name",
    "org": "company_name",
    # domain
    "company_domain": "company_domain",
    "company domain": "company_domain",
    "domain": "company_domain",
    "website": "company_domain",
    # title / role
    "title": "title",
    "job_title": "title",
    "job title": "title",
    "role": "title",
    "position": "title",
    "designation": "title",
    # phone
    "phone": "phone",
    "phone_number": "phone",
    "phone number": "phone",
    "mobile": "phone",
    "telephone": "phone",
    # linkedin
    "linkedin_url": "linkedin_url",
    "linkedin url": "linkedin_url",
    "linkedin": "linkedin_url",
    "linkedin_profile": "linkedin_url",
    "linkedin profile": "linkedin_url",
    # industry
    "industry": "industry",
    # company size
    "company_size": "company_size",
    "company size": "company_size",
    "employees": "company_size",
    "headcount": "company_size",
    "size": "company_size",
}


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

class ProspectFileParser:
    """Unified parser for CSV and XLSX prospect files."""

    # ------------------------------------------------------------------
    # File-level helpers
    # ------------------------------------------------------------------

    @staticmethod
    def detect_format(filename: str) -> str:
        """Return 'csv' or 'xlsx' based on file extension."""
        name = (filename or "").lower()
        if name.endswith(".xlsx"):
            return "xlsx"
        if name.endswith(".csv"):
            return "csv"
        raise FileValidationError(
            f"Unsupported file type: {filename}. Upload a .csv or .xlsx file."
        )

    @staticmethod
    def compute_file_hash(content: bytes) -> str:
        return hashlib.sha256(content).hexdigest()

    @staticmethod
    async def read_bytes(file: UploadFile) -> bytes:
        """Read full file bytes with size validation."""
        content = await file.read()
        if len(content) == 0:
            raise FileValidationError("File is empty")
        if len(content) > MAX_FILE_SIZE:
            raise FileValidationError(
                f"File size ({len(content):,} bytes) exceeds {MAX_FILE_SIZE // (1024*1024)} MB limit"
            )
        await file.seek(0)
        return content

    # ------------------------------------------------------------------
    # Phase 1 — Preview (headers + sample rows)
    # ------------------------------------------------------------------

    @staticmethod
    async def preview(
        file: UploadFile,
    ) -> Dict[str, Any]:
        """Read the file and return headers, sample rows, and auto-mapped columns.

        Returns::

            {
                "format": "csv" | "xlsx",
                "headers": ["Email Address", "First", "Company", ...],
                "sample_rows": [ ["john@acme.com", "John", "Acme"], ... ],
                "auto_mapping": {"Email Address": "email", "First": "first_name", ...},
                "total_rows": 1234,
                "file_size": 54321,
                "file_hash": "abc123...",
            }
        """
        fmt = ProspectFileParser.detect_format(file.filename or "")
        raw = await ProspectFileParser.read_bytes(file)

        if fmt == "xlsx":
            headers, rows, total = ProspectFileParser._read_xlsx(raw)
        else:
            headers, rows, total = ProspectFileParser._read_csv(raw)

        if not headers:
            raise FileValidationError("File has no headers")

        # Auto-map headers using alias table
        auto_mapping: Dict[str, str] = {}
        for h in headers:
            key = h.strip().lower()
            if key in HEADER_ALIASES:
                auto_mapping[h] = HEADER_ALIASES[key]

        return {
            "format": fmt,
            "headers": headers,
            "sample_rows": rows[:PREVIEW_ROWS],
            "auto_mapping": auto_mapping,
            "total_rows": total,
            "file_size": len(raw),
            "file_hash": ProspectFileParser.compute_file_hash(raw),
        }

    # ------------------------------------------------------------------
    # Phase 2 — Full parse with column mapping
    # ------------------------------------------------------------------

    @staticmethod
    async def parse(
        file: UploadFile,
        column_mapping: Dict[str, str],
    ) -> Tuple[List[Dict[str, str]], Dict[str, Any]]:
        """Parse every row using the provided column mapping.

        Args:
            file: The uploaded file (seek(0) will be called).
            column_mapping: Maps file header → system field.
                            e.g. ``{"Email Address": "email", "First": "first_name"}``

        Returns:
            ``(prospects, report)`` where *prospects* is a list of dicts
            with system-field keys and *report* contains stats / errors.
        """
        # Validate mapping has required fields
        mapped_fields = set(column_mapping.values())
        missing = REQUIRED_SYSTEM_FIELDS - mapped_fields
        if missing:
            raise FileValidationError(
                f"Column mapping missing required fields: {', '.join(sorted(missing))}"
            )

        fmt = ProspectFileParser.detect_format(file.filename or "")
        raw = await ProspectFileParser.read_bytes(file)

        if fmt == "xlsx":
            headers, all_rows, total = ProspectFileParser._read_xlsx(raw, max_rows=MAX_ROWS)
        else:
            headers, all_rows, total = ProspectFileParser._read_csv(raw, max_rows=MAX_ROWS)

        # Build header-index → system-field lookup
        header_to_field: Dict[int, str] = {}
        for idx, h in enumerate(headers):
            sys_field = column_mapping.get(h)
            if sys_field and sys_field in SYSTEM_FIELDS:
                header_to_field[idx] = sys_field

        prospects: List[Dict[str, str]] = []
        errors: List[str] = []
        warnings: List[str] = []
        seen_emails: set[str] = set()

        for row_idx, row_values in enumerate(all_rows, start=2):  # row 1 = headers
            record: Dict[str, str] = {f: "" for f in SYSTEM_FIELDS}
            for col_idx, sys_field in header_to_field.items():
                val = row_values[col_idx] if col_idx < len(row_values) else ""
                record[sys_field] = str(val).strip() if val else ""

            email = record["email"].lower()

            # Validate email
            if not email:
                errors.append(f"Row {row_idx}: Missing email")
                continue
            if not EMAIL_RE.match(email):
                errors.append(f"Row {row_idx}: Invalid email: {email}")
                continue

            # Dedup within file
            if email in seen_emails:
                warnings.append(f"Row {row_idx}: Duplicate email: {email}")
                continue
            seen_emails.add(email)

            record["email"] = email

            # Soft validation
            if record["linkedin_url"] and not record["linkedin_url"].startswith("http"):
                warnings.append(
                    f"Row {row_idx}: LinkedIn URL should start with http: {record['linkedin_url']}"
                )

            prospects.append(record)

            if len(prospects) >= MAX_ROWS:
                errors.append(f"File exceeds maximum {MAX_ROWS:,} rows")
                break

        report = {
            "total_rows": total,
            "valid_prospects": len(prospects),
            "errors": errors,
            "warnings": warnings,
            "duplicate_emails": list(seen_emails & {w.split(": ")[-1] for w in warnings if "Duplicate" in w}),
            "headers_found": headers,
            "is_valid": len(errors) == 0,
        }

        return prospects, report

    # ------------------------------------------------------------------
    # Legacy compat — parse_and_validate (old CSV-only path)
    # ------------------------------------------------------------------

    @staticmethod
    async def parse_and_validate(
        file: UploadFile,
        validate_only: bool = False,
    ) -> Tuple[List[Dict[str, str]], Dict[str, Any]]:
        """Backwards-compatible entry point used by the old upload endpoint.

        Auto-maps headers using the alias table so files with standard
        column names work without an explicit mapping step.
        """
        fmt = ProspectFileParser.detect_format(file.filename or "")
        raw = await ProspectFileParser.read_bytes(file)

        limit = 100 if validate_only else MAX_ROWS
        if fmt == "xlsx":
            headers, all_rows, total = ProspectFileParser._read_xlsx(raw, max_rows=limit)
        else:
            headers, all_rows, total = ProspectFileParser._read_csv(raw, max_rows=limit)

        if not headers:
            raise FileValidationError("File has no headers")

        # Auto-map using alias table
        mapping: Dict[str, str] = {}
        for h in headers:
            key = h.strip().lower()
            if key in HEADER_ALIASES:
                mapping[h] = HEADER_ALIASES[key]

        if "email" not in set(mapping.values()):
            raise FileValidationError(
                "Could not find an 'email' column. "
                f"Headers found: {', '.join(headers)}"
            )

        # Re-seek and delegate to parse()
        await file.seek(0)
        return await ProspectFileParser.parse(file, mapping)

    # ------------------------------------------------------------------
    # Dedup helper (static, used by process endpoint)
    # ------------------------------------------------------------------

    @staticmethod
    async def deduplicate_prospects(
        prospects: List[Dict],
    ) -> Tuple[List[Dict], List[str]]:
        seen: Dict[str, bool] = {}
        unique: List[Dict] = []
        dupes: List[str] = []
        for p in prospects:
            email = p["email"].lower()
            if email not in seen:
                seen[email] = True
                unique.append(p)
            else:
                dupes.append(email)
        return unique, dupes

    # ------------------------------------------------------------------
    # Internal readers
    # ------------------------------------------------------------------

    @staticmethod
    def _read_csv(
        raw: bytes,
        max_rows: int = MAX_ROWS,
    ) -> Tuple[List[str], List[List[str]], int]:
        """Read CSV bytes → (headers, rows, total_row_count)."""
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = raw.decode("latin-1")
            except UnicodeDecodeError:
                raise FileValidationError("File encoding not recognised (expected UTF-8)")

        reader = csv.reader(io.StringIO(text))
        rows_iter = iter(reader)

        try:
            headers = next(rows_iter)
        except StopIteration:
            return [], [], 0

        headers = [h.strip() for h in headers]

        rows: List[List[str]] = []
        total = 0
        for row in rows_iter:
            total += 1
            if len(rows) < max_rows:
                rows.append(row)

        return headers, rows, total

    @staticmethod
    def _read_xlsx(
        raw: bytes,
        max_rows: int = MAX_ROWS,
    ) -> Tuple[List[str], List[List[str]], int]:
        """Read XLSX bytes → (headers, rows, total_row_count)."""
        if not XLSX_AVAILABLE:
            raise FileValidationError(
                "XLSX support not available — openpyxl is not installed"
            )

        wb = load_workbook(filename=io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        if ws is None:
            raise FileValidationError("XLSX file has no active worksheet")

        rows_iter = ws.iter_rows(values_only=True)

        # First row = headers
        try:
            header_row = next(rows_iter)
        except StopIteration:
            wb.close()
            return [], [], 0

        headers = [str(c).strip() if c is not None else "" for c in header_row]
        # Drop trailing empty headers
        while headers and not headers[-1]:
            headers.pop()

        rows: List[List[str]] = []
        total = 0
        for row in rows_iter:
            # Skip fully empty rows
            if all(c is None or str(c).strip() == "" for c in row):
                continue
            total += 1
            if len(rows) < max_rows:
                rows.append([str(c).strip() if c is not None else "" for c in row])

        wb.close()
        return headers, rows, total


# ---------------------------------------------------------------------------
# Backwards-compat alias — old code imports ProspectCSVParser from csv_parser
# ---------------------------------------------------------------------------

ProspectCSVParser = ProspectFileParser
